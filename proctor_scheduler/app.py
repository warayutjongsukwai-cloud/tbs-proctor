#!/usr/bin/env python3
"""UI ระบบจัดเจ้าหน้าที่คุมสอบ TBS (Streamlit).

รัน:  streamlit run app.py
Flow: อัปโหลดไฟล์ -> ปรับ demand/คาบ -> ทำเครื่องหมายวันไม่ว่าง
      -> จัดคิวอัตโนมัติ -> ตรวจ/แก้มือ -> ดาวน์โหลด Excel
"""
from __future__ import annotations
import statistics
import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from proctor_scheduler.models import (Problem, Staff, Session, ExamType, PayTier,
                                      Unavailability, Assignment, PAY_YELLOW, PAY_GREEN)
from proctor_scheduler.demand import estimate_sessions
from proctor_scheduler.planner import plan_rooms
from proctor_scheduler.staff_io import load_roster
from proctor_scheduler.engine import solve, feasibility_check, detect_conflicts
from proctor_scheduler.excel_io import export

st.set_page_config(page_title="จัดคนคุมสอบ TBS", page_icon="📋", layout="wide")
st.title("📋 ระบบจัดเจ้าหน้าที่คุมสอบ TBS")
st.caption("เกลี่ยเงิน/คิวให้แฟร์ด้วย CP-SAT · เจ้าหน้าที่ตรวจและแก้ได้ก่อนใช้จริง")

ss = st.session_state
ss.setdefault("sessions", None)
ss.setdefault("staff", None)
ss.setdefault("unavail", {})
ss.setdefault("result", None)
ss.setdefault("locks", set())        # {(staff_id, session_id)} บังคับจัด (manual override)
ss.setdefault("allocations", None)   # แผนจัดห้อง (ถ้าเลือกโหมดจัดห้อง)


def _save_tmp(uploaded) -> str:
    path = f"/tmp/{uploaded.name}"
    with open(path, "wb") as f:
        f.write(uploaded.getbuffer())
    return path


# ======== แถบซ้าย: อัปโหลด ========
with st.sidebar:
    st.header("1️⃣ ไฟล์ข้อมูล")
    mode = st.radio("วิธีได้จำนวนที่ต้องการ", [
        "จัดห้องอัตโนมัติ (พบ.>SC1>SC3)",
        "ประเมินจาก section",
        "จากชีทคิว (เช่น ท่าพระจันทร์)"],
        help="จัดห้อง/ประเมิน = ใช้ตารางสอนรังสิต · จากชีทคิว = อ่านจำนวนที่ต้องการตรงจากชีทคิว")
    need_sched = not mode.startswith("จากชีทคิว")
    f_sched = st.file_uploader("ตารางสอน (.xlsx)", type="xlsx") if need_sched else None
    f_roster = st.file_uploader("ไฟล์คิว (ดึงรายชื่อ/จำนวนที่ต้องการ)", type="xlsx")
    exam_type = st.radio("ประเภทสอบ", ["กลางภาค", "ปลายภาค"], horizontal=True)
    et = ExamType.MIDTERM if exam_type == "กลางภาค" else ExamType.FINAL
    campus = st.text_input("ศูนย์", "รังสิต")
    term = st.text_input("ภาคการศึกษา", "1/2569")
    default_proc = st.number_input("คนคุม/ห้อง (โหมดประเมินเท่านั้น)", 1, 5, 2)

    roster_sheet = None
    if f_roster:
        roster_sheet = st.selectbox(
            "ชีทรายชื่อ/คิว", load_workbook(_save_tmp(f_roster), read_only=True).sheetnames)
    roster_first_row = st.number_input("แถวแรกของรายชื่อ", 2, 20, 7,
                                       help="รังสิต=7, ท่าพระจันทร์=6 (ปรับตามไฟล์)")

    st.divider()
    st.caption("➕ ยอดยกมา (ไม่บังคับ) — เกลี่ยเงินต่อเนื่องข้ามเทอม")
    f_carry = st.file_uploader("ไฟล์คิวเทอมก่อน", type="xlsx", key="fc")
    carry_sheet = None
    if f_carry:
        carry_sheet = st.selectbox(
            "ชีทเทอมก่อน", load_workbook(_save_tmp(f_carry), read_only=True).sheetnames)

    ready = bool(f_roster) and (bool(f_sched) or not need_sched)
    if st.button("➡️ เริ่ม: อ่านไฟล์", type="primary", disabled=not ready):
        from proctor_scheduler.demand import build_sessions_from_queue
        if mode.startswith("จัดห้อง"):
            rp = plan_rooms(_save_tmp(f_sched), et, campus)
            ss.sessions, ss.allocations = rp.sessions, rp.allocations
            if rp.unplaced:
                st.warning(f"⚠️ {len(rp.unplaced)} คาบห้องไม่พอ (ดูแผนห้อง)")
        elif mode.startswith("ประเมิน"):
            ss.sessions = estimate_sessions(_save_tmp(f_sched), et, int(default_proc))
            ss.allocations = None
        else:   # จากชีทคิว
            ss.sessions = build_sessions_from_queue(
                _save_tmp(f_roster), roster_sheet, et, int(roster_first_row))
            ss.allocations = None
        ss.staff = load_roster(_save_tmp(f_roster), roster_sheet, campus,
                               first_row=int(roster_first_row))
        ss.unavail = {s.staff_id: set() for s in ss.staff}
        ss.result = None
        msg = f"อ่านแล้ว: {len(ss.sessions)} คาบ · {len(ss.staff)} คน"
        if f_carry and carry_sheet:
            from proctor_scheduler.staff_io import load_carry
            n = load_carry(_save_tmp(f_carry), carry_sheet, ss.staff)
            msg += f" · ยอดยกมา {n} คน"
        if ss.allocations:
            msg += " · จัดห้องแล้ว"
        st.success(msg)

if ss.sessions is None:
    st.info("⬅️ อัปโหลดตารางสอน + ไฟล์คิว ที่แถบซ้าย แล้วกด 'เริ่ม'")
    st.stop()

sessions: list[Session] = ss.sessions
staff: list[Staff] = ss.staff
avail_staff = [s for s in staff if not s.is_exempt]

tab1, tab2, tab3, tab4 = st.tabs(
    ["2️⃣ ปรับจำนวนที่ต้องการ", "3️⃣ วันไม่ว่าง", "4️⃣ จัดคิว", "5️⃣ ผล & ดาวน์โหลด"])

# ======== แท็บ 2: ปรับ demand ========
with tab1:
    st.subheader("จำนวนคนคุมที่ต้องการต่อคาบ (แก้ได้)")
    st.caption("ค่านี้ประเมินจากจำนวน section — ปรับตามคนสอบจริง / เพิ่มคาบชดเชยได้")
    df = pd.DataFrame([{
        "คาบ": f"{s.date} {s.slot}", "เวลา": f"{s.start_time}-{s.end_time}",
        "ประเภท": "🟡เหลือง(300)" if s.tier == PayTier.YELLOW else "🟢เขียว(100)",
        "คนคุมที่ต้องการ": s.demand} for s in sessions])
    edited = st.data_editor(df, hide_index=True, use_container_width=True,
                            disabled=["คาบ", "เวลา", "ประเภท"])
    for i, s in enumerate(sessions):
        s.demand = int(edited.iloc[i]["คนคุมที่ต้องการ"])
    c1, c2, c3 = st.columns(3)
    c1.metric("คาบทั้งหมด", len(sessions))
    c2.metric("ต้องการรวม (คน-คาบ)", sum(s.demand for s in sessions))
    c3.metric("เฉลี่ยคิว/คน", f"{sum(s.demand for s in sessions)/max(len(avail_staff),1):.1f}")

# ======== แท็บ 3: วันไม่ว่าง ========
with tab2:
    st.subheader("ทำเครื่องหมายคนที่ไม่ว่าง / ขอยกเว้น")

    # ---- โหมดเพิ่มหลายคนทีเดียว (เช่น หลายคนลาคาบเดียวกัน) ----
    st.markdown("**➕ เพิ่มหลายคนทีเดียว** — เลือกคนหลายคน + คาบที่ทุกคนคุมไม่ได้ แล้วกดเพิ่ม")
    opts_all = {f"{s.date} {s.slot}": s.session_id for s in sessions}
    bcol1, bcol2 = st.columns(2)
    with bcol1:
        bulk_people = st.multiselect("เจ้าหน้าที่ (เลือกได้หลายคน)", [s.name for s in staff])
    with bcol2:
        bulk_slots = st.multiselect("คาบที่ไม่ว่าง (เลือกได้หลายคาบ)", list(opts_all.keys()))
    bc1, bc2 = st.columns(2)
    if bc1.button("➕ เพิ่มเป็นไม่ว่าง", disabled=not (bulk_people and bulk_slots)):
        sids = {opts_all[k] for k in bulk_slots}
        for nm in bulk_people:
            stf = next(s for s in staff if s.name == nm)
            ss.unavail.setdefault(stf.staff_id, set()).update(sids)
        st.success(f"เพิ่มแล้ว: {len(bulk_people)} คน × {len(bulk_slots)} คาบ")
    if bc2.button("➖ เอาออกจากไม่ว่าง", disabled=not (bulk_people and bulk_slots)):
        sids = {opts_all[k] for k in bulk_slots}
        for nm in bulk_people:
            stf = next(s for s in staff if s.name == nm)
            ss.unavail.get(stf.staff_id, set()).difference_update(sids)
        st.success(f"เอาออกแล้ว: {len(bulk_people)} คน × {len(bulk_slots)} คาบ")

    st.divider()
    # ---- โหมดรายคน (ละเอียด + ยกเว้นทั้งเทอม) ----
    st.markdown("**🔎 แก้รายคน** — เลือกคนเดียวเพื่อดู/ปรับละเอียด หรือยกเว้นทั้งเทอม")
    colL, colR = st.columns([1, 2])
    with colL:
        pick = st.selectbox("เลือกเจ้าหน้าที่", [s.name for s in staff])
        picked = next(s for s in staff if s.name == pick)
        picked.is_exempt = st.checkbox("ยกเว้นทั้งเทอม", picked.is_exempt)
    with colR:
        if picked.is_exempt:
            st.info(f"{pick} = ยกเว้นทั้งเทอม (ไม่ถูกจัดเลย)")
        else:
            cur = ss.unavail.get(picked.staff_id, set())
            chosen = st.multiselect("คาบที่ไม่ว่าง", list(opts_all.keys()),
                                    default=[k for k, v in opts_all.items() if v in cur])
            ss.unavail[picked.staff_id] = {opts_all[k] for k in chosen}

    # ---- สรุปรายคนที่มีวันไม่ว่าง ----
    busy = [(next(s.name for s in staff if s.staff_id == sid), len(v))
            for sid, v in ss.unavail.items() if v]
    if busy:
        st.caption("คนที่มีวันไม่ว่าง: " + " · ".join(f"{n}({c})" for n, c in sorted(busy)))
    st.write(f"รวม: ไม่ว่าง {sum(len(v) for v in ss.unavail.values())} ช่อง · "
             f"ยกเว้นทั้งเทอม {sum(1 for s in staff if s.is_exempt)} คน")

# ======== แท็บ 4: จัดคิว ========
with tab3:
    st.subheader("จัดคิวอัตโนมัติ")
    time_limit = st.slider("เวลาค้นหาสูงสุด (วินาที)", 5, 120, 30)

    def _build_problem():
        return Problem(
            staff=staff, sessions=sessions, campus=campus, term=term,
            unavailable=[Unavailability(sid, se)
                         for sid, v in ss.unavail.items() for se in v],
            locked=[Assignment(a, b) for (a, b) in ss.locks])

    problem = _build_problem()
    ok, msg = feasibility_check(problem)
    (st.success if ok else st.warning)(msg)
    allow_partial = st.checkbox(
        "จัดเท่าที่ทำได้ ถ้าคนไม่พอ (ไม่ต้องครบทุกคาบ)", value=not ok,
        help="เปิดไว้ = ถึงคนไม่พอก็จัดให้มากสุดเท่าที่ทำได้ แล้วบอกว่าคาบไหนยังขาดกี่คน")
    if ss.locks:
        st.info(f"🔒 มีการล็อกไว้ {len(ss.locks)} รายการ (จากแท็บผล) — จะคงไว้ตอนจัด")
    if st.button("🚀 จัดคิว", type="primary"):
        with st.spinner("กำลังจัด..."):
            ss.result = solve(problem, time_limit_s=time_limit, allow_partial=allow_partial)
        r = ss.result
        if r.is_ok():
            pays = [r.pay[s.staff_id] for s in avail_staff]
            st.success(f"สำเร็จ ({r.status}, {r.solve_seconds:.1f}s)")
            m = st.columns(4)
            m[0].metric("เงิน ต่ำสุด–สูงสุด", f"{min(pays)}–{max(pays)}")
            m[1].metric("ส่วนต่างเงิน", max(pays) - min(pays))
            m[2].metric("SD เงิน", f"{statistics.pstdev(pays):.0f}")
            m[3].metric("ช่องว่างกลางวัน", r.gaps, help="ยิ่งน้อยยิ่งดี (คาบติดกัน)")
            if r.shortfalls:
                sess_map = {s.session_id: s for s in sessions}
                total_miss = sum(r.shortfalls.values())
                lines = "\n".join(
                    f"• {sess_map[sid].date} {sess_map[sid].slot} — ขาดอีก {n} คน"
                    for sid, n in sorted(r.shortfalls.items()))
                st.warning(f"⚠️ จัดเท่าที่ทำได้: ยังขาดรวม {total_miss} คน-คาบ "
                           f"({len(r.shortfalls)} คาบไม่เต็ม)\n\n{lines}")
            else:
                st.info("✅ จัดครบทุกคาบ ไม่มีคาบไหนขาดคน")
        else:
            st.error(f"จัดไม่ได้ ({r.status}) — ลองเปิด 'จัดเท่าที่ทำได้' ด้านบน "
                     f"หรือดูแท็บ 'ต้องการ'/วันไม่ว่าง")

# ======== แท็บ 5: ผล + ดาวน์โหลด ========
with tab4:
    st.subheader("ผลการจัด & ดาวน์โหลด")
    r = ss.result
    if r is None or not r.is_ok():
        st.info("ยังไม่มีผล — ไปที่แท็บ 'จัดคิว' แล้วกดจัดก่อน")
    else:
        rows = []
        for i, s in enumerate(avail_staff, 1):
            sids = set(r.assignments.get(s.staff_id, []))
            y = sum(1 for se in sessions if se.session_id in sids and se.tier == PayTier.YELLOW)
            g = sum(1 for se in sessions if se.session_id in sids and se.tier == PayTier.GREEN)
            rows.append({"No.": i, "ชื่อ": s.name, "เหลือง": y, "เขียว": g,
                         "คิวรวม": y + g, "เงิน": y * PAY_YELLOW + g * PAY_GREEN})
        dfr = pd.DataFrame(rows)
        mean = dfr["เงิน"].mean()
        st.dataframe(
            dfr.style.apply(
                lambda row: ['background-color:#FFC7CE'
                             if (col == "เงิน" and abs(row["เงิน"] - mean) > 0.15 * mean)
                             else '' for col in dfr.columns], axis=1),
            hide_index=True, use_container_width=True, height=400)
        st.caption("แถวแดง = เงินเบี่ยงจากเฉลี่ย >15% (ควรตรวจ)")

        # ---- Manual Override ----
        with st.expander("🔧 แก้มือ: ล็อกคน–คาบ แล้วจัดใหม่ (ไม่ใช่ black box)"):
            st.caption("ล็อก = บังคับให้คนนี้ได้คาบนี้แน่นอน แล้วกดจัดใหม่ "
                       "ระบบจะเกลี่ยคนอื่นรอบๆ ให้")
            oc1, oc2, oc3 = st.columns([2, 2, 1])
            pname = oc1.selectbox("เจ้าหน้าที่", [s.name for s in avail_staff], key="ov_s")
            sopts = {f"{se.date} {se.slot}": se.session_id for se in sessions}
            psess = oc2.selectbox("คาบ", list(sopts.keys()), key="ov_se")
            pobj = next(s for s in avail_staff if s.name == pname)
            if oc3.button("🔒 ล็อก"):
                ss.locks.add((pobj.staff_id, sopts[psess]))
            if ss.locks:
                st.write("**รายการที่ล็อก:**")
                id2name = {s.staff_id: s.name for s in staff}
                id2sess = {se.session_id: f"{se.date} {se.slot}" for se in sessions}
                for (sidd, seid) in sorted(ss.locks):
                    lc1, lc2 = st.columns([4, 1])
                    lc1.write(f"🔒 {id2name.get(sidd, sidd)} → {id2sess.get(seid, seid)}")
                    if lc2.button("ปลด", key=f"unlock_{sidd}_{seid}"):
                        ss.locks.discard((sidd, seid)); st.rerun()
                if st.button("🔄 จัดใหม่ (คงรายการที่ล็อก)", type="primary"):
                    prob2 = Problem(
                        staff=staff, sessions=sessions, campus=campus, term=term,
                        unavailable=[Unavailability(a, b)
                                     for a, v in ss.unavail.items() for b in v],
                        locked=[Assignment(a, b) for (a, b) in ss.locks])
                    with st.spinner("จัดใหม่..."):
                        ss.result = solve(prob2, time_limit_s=30)
                    st.rerun()

            # ตรวจ conflict ของผลปัจจุบัน
            confs = detect_conflicts(problem, r.assignments)
            if confs:
                st.error("พบ conflict:\n\n" + "\n\n".join(confs[:10]))
            else:
                st.success("✅ ไม่พบ conflict (ไม่ชนวันไม่ว่าง · ทุกคาบคนครบ)")

        if ss.allocations:
            from proctor_scheduler.planner import distribute_proctors, sc_balance_report
            with st.expander("🏫 แผนจัดห้อง (พบ. > SC1 > SC3) + สมดุล SC"):
                id2name = {s.staff_id: s.name for s in staff}
                sess_map = {s.session_id: s for s in sessions}
                dist = distribute_proctors(sessions, ss.allocations, r.assignments)
                plan_rows = []
                for sid in sorted(ss.allocations.keys()):
                    se = sess_map[sid]
                    for a, names in dist.get(sid, []):
                        plan_rows.append({"วันที่": se.date, "ช่วง": se.slot,
                                          "ตึก": a.room.building, "ห้อง": a.room.room_id,
                                          "วิชา": ", ".join(a.courses), "คนสอบ": a.students,
                                          "คนคุม": a.proctors,
                                          "ผู้คุม": ", ".join(id2name.get(n, n) for n in names)})
                st.dataframe(pd.DataFrame(plan_rows), hide_index=True,
                             use_container_width=True, height=360)
                scb = sc_balance_report(dist, id2name)
                if scb:
                    vals = list(scb.values())
                    st.caption(f"สมดุล SC: ไป SC คนละ {min(vals)}–{max(vals)} ครั้ง "
                               f"(เกลี่ยแล้ว — A8)")

        path = f"/tmp/ผลจัดคิว_{campus}_{term.replace('/','-')}.xlsx"
        export(problem, r, path, title=f"คิวกรรมการคุมสอบ {campus} {term}",
               allocations=ss.allocations)
        nsheet = 3 if ss.allocations else 2
        with open(path, "rb") as f:
            st.download_button(f"⬇️ ดาวน์โหลด Excel ({nsheet} ชีท)", f.read(),
                               file_name=path.split("/")[-1], type="primary",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        st.caption("💡 แก้มือได้ในไฟล์ Excel ที่ดาวน์โหลด · เก็บ version เดิมไว้เสมอ ทำ v2 ใหม่")
