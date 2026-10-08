"""โหลดรายชื่อเจ้าหน้าที่ + ยอดสะสม จากไฟล์คิวเดิม."""
from __future__ import annotations
from openpyxl import load_workbook
from .models import Staff


def load_roster(xlsx_path: str, sheet: str, campus: str = "rangsit",
                name_col: int = 2, first_row: int = 7,
                exempt_names: set[str] | None = None) -> list[Staff]:
    """อ่านชื่อจากชีทคิว (คอลัมน์ 'ชื่อ'). ยอดยกมา = 0 (ต้นปี) เว้นระบุเพิ่ม."""
    exempt_names = exempt_names or set()
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet]
    staff = []
    for r in range(first_row, ws.max_row + 1):
        no = ws.cell(r, 1).value
        nm = ws.cell(r, name_col).value
        if not (isinstance(no, (int, float)) and nm):
            continue
        nm = str(nm).strip()
        staff.append(Staff(
            staff_id=f"ST{int(no):03d}", name=nm, full_name=nm, campus=campus,
            is_exempt=nm in exempt_names))
    return staff


def load_carry(xlsx_path: str, sheet: str, staff: list, name_col: int = 2) -> int:
    """อ่าน "ยอดสะสมทั้งหมด" จากชีทเทอมก่อน ใส่เข้า Staff.carry_* (A2).

    หาคอลัมน์ "รวมทั้งหมด" อัตโนมัติ = คู่ (คิว, เงิน) ขวาสุดในแถวหัว คิว/เงิน
    (layout ต่างกันแต่ละเทอม จึงไม่ fix ตำแหน่ง). คืนจำนวนคนที่แมตช์ได้.
    หมายเหตุ: ชีทเก่าไม่แยกเหลือง/เขียว *สะสม* → เก็บเป็น carry_pay (เงินรวม)
    + carry_green (คิวรวม เป็น proxy ภาระสะสม); carry_yellow=0.
    """
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet]
    # หาแถวหัวที่เป็น คิว/เงิน สลับกัน
    hdr_row = None
    for r in range(1, 12):
        vals = [str(ws.cell(r, c).value or '') for c in range(1, ws.max_column + 1)]
        if vals.count('คิว') >= 2 and vals.count('เงิน') >= 2:
            hdr_row = r
            break
    if hdr_row is None:
        return 0
    p_cols = [c for c in range(1, ws.max_column + 1) if ws.cell(hdr_row, c).value == 'เงิน']
    # บล็อกสรุปเป็นคู่ (คิว,เงิน) ติดกัน ส่วน grid คาบมีแต่ 'คิว' เดี่ยว ๆ
    # -> เงินรวม = คอลัมน์ 'เงิน' ขวาสุด, คิวรวม = ช่องซ้ายของมัน
    tot_p = p_cols[-1]
    tot_q = tot_p - 1
    first_row = hdr_row + 1
    by_name = {s.name: s for s in staff}
    matched = 0
    for r in range(first_row, ws.max_row + 1):
        nm = ws.cell(r, name_col).value
        if not nm:
            continue
        s = by_name.get(str(nm).strip())
        if s is None:
            continue
        q = ws.cell(r, tot_q).value or 0
        p = ws.cell(r, tot_p).value or 0
        try:
            s.carry_green = int(q); s.carry_yellow = 0; s.carry_pay = int(p)
            matched += 1
        except (TypeError, ValueError):
            continue
    return matched
