"""สร้างไฟล์ 'ผลจัดคิว' ให้ตรงฟอร์แมตต้นฉบับของคณะ (เมทริกซ์ คน × คาบ).

โครงตามไฟล์จริง 'คิวจัดจนท.สอบรังสิต':
  แถว1  ชื่อเรื่อง
  แถว2  No. | ชื่อผู้คุมสอบ | ชื่อ-สกุล | [คิวสอบ<ภาค>] (merge D-I) | <วันสอบ> (merge ต่อวัน)
  แถว3  เหลือง | เขียว | รวม | <ช่วง: เช้า/เที่ยง/เย็น>
  แถว4  <เวลา เช่น 9-11.30>
  แถว5  คิว | เงิน (×3) | คิว (ทุกคาบ)
  แถว6+ ข้อมูล: ใส่ 1 ช่องที่ถูกจัด, D-I เป็นสูตรรวม
"""
from __future__ import annotations
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from .models import PayTier, PAY_YELLOW, PAY_GREEN

_TH_MONTH = {"ม.ค.": 1, "ก.พ.": 2, "มี.ค.": 3, "เม.ย.": 4, "พ.ค.": 5, "มิ.ย.": 6,
             "ก.ค.": 7, "ส.ค.": 8, "ก.ย.": 9, "ต.ค.": 10, "พ.ย.": 11, "ธ.ค.": 12}
_SLOT_ORDER = {"เช้า": 0, "เที่ยง": 1, "บ่าย": 1, "เย็น": 2}

_BORDER = Border(*[Side(style="thin", color="D0D0D0")] * 4)
_HEAD_FILL = PatternFill("solid", fgColor="E7EAF3")
_YEL_FILL = PatternFill("solid", fgColor="FFF2CC")
_GRN_FILL = PatternFill("solid", fgColor="E2EFDA")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_BOLD = Font(bold=True)


def _date_key(date_str: str):
    """'จ.21 ก.ย.2569' -> (ปี, เดือน, วัน) สำหรับเรียงตามจริง."""
    parts = date_str.split()
    day = int(re.search(r"\d+", parts[0]).group())
    m = re.match(r"([ก-ฮ.]+)(\d+)", parts[1]) if len(parts) > 1 else None
    month = _TH_MONTH.get(m.group(1), 0) if m else 0
    year = int(m.group(2)) if m else 0
    return (year, month, day)


def build_queue_file(sessions, staff, assignments: dict, path: str,
                     title: str = "คิวกรรมการคุมสอบ ศูนย์รังสิต",
                     term_label: str = "คิวสอบ") -> str:
    # เรียงคาบตามวันจริง แล้วช่วงเวลา
    sess = sorted(sessions, key=lambda s: (_date_key(s.date), _SLOT_ORDER.get(s.slot, 9)))
    wb = Workbook()
    ws = wb.active
    ws.title = "ผลจัดคิว"

    C_NO, C_NICK, C_FULL = 1, 2, 3
    C_SUM0 = 4                       # D: เหลืองคิว E เหลืองเงิน F เขียวคิว G เขียวเงิน H รวมคิว I รวมเงิน
    C_SESS0 = 10                     # J เป็นต้นไป = คาบสอบ
    n_sess = len(sess)
    last_col = C_SESS0 + n_sess - 1

    # แถว 1: ชื่อเรื่อง
    ws.cell(1, 1, title)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")

    # แถว 2-5 หัวซ้าย
    ws.cell(2, C_NO, "No."); ws.cell(2, C_NICK, "ชื่อผู้คุมสอบ"); ws.cell(2, C_FULL, "ชื่อ-สกุล")
    for r in (C_NO, C_NICK, C_FULL):
        ws.merge_cells(start_row=2, start_column=r, end_row=5, end_column=r)
    ws.cell(2, C_SUM0, term_label)
    ws.merge_cells(start_row=2, start_column=C_SUM0, end_row=2, end_column=C_SUM0 + 5)
    for (c, lab) in ((C_SUM0, "เหลือง"), (C_SUM0 + 2, "เขียว"), (C_SUM0 + 4, "รวม")):
        ws.cell(3, c, lab); ws.merge_cells(start_row=3, start_column=c, end_row=3, end_column=c + 1)
    for i, lab in enumerate(["คิว", "เงิน"] * 3):
        ws.cell(5, C_SUM0 + i, lab)

    # แถว 2-5 หัวคาบ (group ตามวัน -> merge แถว 2)
    yellow_cols, green_cols = [], []
    col = C_SESS0
    i = 0
    while i < n_sess:
        day = sess[i].date
        j = i
        while j < n_sess and sess[j].date == day:
            j += 1
        c0, c1 = col, col + (j - i) - 1
        ws.cell(2, c0, day)
        if c1 > c0:
            ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c1)
        for k in range(i, j):
            cc = C_SESS0 + k
            s = sess[k]
            ws.cell(3, cc, s.slot)
            ws.cell(4, cc, f"{s.start_time}-{s.end_time}")
            ws.cell(5, cc, "คิว")
            (yellow_cols if s.tier == PayTier.YELLOW else green_cols).append(get_column_letter(cc))
        col = c1 + 1
        i = j

    # แถวข้อมูล
    id2idx = {s.staff_id: set(assignments.get(s.staff_id, [])) for s in staff}
    sid_order = [s.session_id for s in sess]
    r = 6
    for no, st in enumerate(staff, 1):
        ws.cell(r, C_NO, no)
        ws.cell(r, C_NICK, st.name)
        ws.cell(r, C_FULL, getattr(st, "full_name", "") or "")
        got = id2idx[st.staff_id]
        for k, sid in enumerate(sid_order):
            if sid in got:
                ws.cell(r, C_SESS0 + k, 1)
        # สูตรสรุป (เหมือนต้นฉบับ)
        yq = "+".join(f"{c}{r}" for c in yellow_cols) or "0"
        gq = "+".join(f"{c}{r}" for c in green_cols) or "0"
        ws.cell(r, C_SUM0, f"={yq}")                     # เหลือง คิว
        ws.cell(r, C_SUM0 + 1, f"={get_column_letter(C_SUM0)}{r}*{PAY_YELLOW}")   # เหลือง เงิน
        ws.cell(r, C_SUM0 + 2, f"={gq}")                 # เขียว คิว
        ws.cell(r, C_SUM0 + 3, f"={get_column_letter(C_SUM0+2)}{r}*{PAY_GREEN}")  # เขียว เงิน
        ws.cell(r, C_SUM0 + 4, f"={get_column_letter(C_SUM0)}{r}+{get_column_letter(C_SUM0+2)}{r}")       # รวม คิว
        ws.cell(r, C_SUM0 + 5, f"={get_column_letter(C_SUM0+1)}{r}+{get_column_letter(C_SUM0+3)}{r}")     # รวม เงิน
        r += 1

    # จัดรูป: หัวตาราง + เส้น + สี
    for rr in range(2, 6):
        for cc in range(1, last_col + 1):
            cell = ws.cell(rr, cc)
            cell.font = _BOLD; cell.alignment = _CENTER; cell.border = _BORDER
            cell.fill = _HEAD_FILL
    # สีเหลือง/เขียวบนหัวคอลัมน์สรุป
    for cc in (C_SUM0, C_SUM0 + 1):
        for rr in range(3, 6): ws.cell(rr, cc).fill = _YEL_FILL
    for cc in (C_SUM0 + 2, C_SUM0 + 3):
        for rr in range(3, 6): ws.cell(rr, cc).fill = _GRN_FILL
    for rr in range(6, r):
        for cc in range(1, last_col + 1):
            cell = ws.cell(rr, cc); cell.border = _BORDER
            if cc >= C_SUM0: cell.alignment = _CENTER
    # ความกว้าง
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 24
    for cc in range(C_SUM0, last_col + 1):
        ws.column_dimensions[get_column_letter(cc)].width = 7
    ws.freeze_panes = "D6"
    wb.save(path)
    return path
