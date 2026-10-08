"""อ่านตารางสอน -> รายวิชา/section ต่อคาบ สำหรับจัดห้อง.

กฎ: ต่าง section = แยกห้อง · วิชาคนสอบ <10 คน รวมกันในห้องเดียวได้ (ผู้ใช้ 8 ต.ค.)
หมายเหตุ: ตารางสอนมีแค่ "ที่นั่งที่เปิด" ไม่ใช่คนสอบจริง -> students เป็นค่าประมาณ
"""
from __future__ import annotations
import re
from dataclasses import dataclass
from openpyxl import load_workbook
from .demand import (FACULTY_SHEETS, _scan_sheet, _parse_exam_cell,
                     _slot_from_time, _enroll, _is_weekend)
from .models import ExamType, ROOM_STUDENT_CAP

SMALL_CLASS = 10          # วิชาคนสอบ < 10 รวมกันได้


@dataclass
class ExamUnit:
    """1 หน่วยสอบ = 1 ห้องที่จะจัด (section เดี่ยว หรือกลุ่มวิชาเล็กรวมกัน)."""
    label: str            # เช่น "AC201(01)" หรือ "รวมเล็ก: X,Y,Z"
    students: int


def exams_per_session(xlsx_path: str, exam_type: ExamType = ExamType.MIDTERM
                      ) -> dict[tuple, list[ExamUnit]]:
    """คืน dict (date,slot,start,end) -> list[ExamUnit] (จัด section + split>65 + รวมเล็ก)."""
    which = 'mid' if exam_type == ExamType.MIDTERM else 'final'
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    raw: dict[tuple, list[tuple[str, int]]] = {}
    for name in FACULTY_SHEETS:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        # ต้องดึง "รหัสวิชา+sec" ด้วย -> อ่านเองจาก sheet
        rows = list(ws.iter_rows(values_only=True))
        hdr = next((i for i, r in enumerate(rows)
                    if 'วันสอบกลางภาค' in ' '.join(str(v) for v in r if v)
                    and 'จำนวนนักศึกษา' in ' '.join(str(v) for v in r if v)), None)
        if hdr is None:
            continue
        col = {}
        for k, v in enumerate(rows[hdr]):
            if v is None:
                continue
            sv = str(v)
            if 'วันสอบกลางภาค' in sv: col['mid'] = k
            elif 'วันสอบปลายภาค' in sv: col['final'] = k
            elif 'จำนวนนักศึกษา' in sv: col['enr'] = k
            elif sv == 'Sec': col['sec'] = k
        want = col.get(which)
        if want is None:
            continue
        for r in rows[hdr + 2:]:
            code = r[0] if r else None
            if not code or not re.match(r'^[A-Z]{2,4}\d', str(code).strip()):
                continue
            parsed = _parse_exam_cell(r[want] if want < len(r) else None)
            if parsed is None:
                continue
            date, start, end = parsed
            slot = _slot_from_time(start + '-' + end if end else start)
            sec = str(r[col['sec']]).strip() if 'sec' in col and col['sec'] < len(r) and r[col['sec']] else ''
            enr = _enroll(r[col['enr']]) if 'enr' in col and col['enr'] < len(r) else 0
            if enr <= 0:
                continue
            key = (date, slot, start, end)
            raw.setdefault(key, []).append((f"{str(code).strip()}({sec})", enr))

    out: dict[tuple, list[ExamUnit]] = {}
    for key, items in raw.items():
        units: list[ExamUnit] = []
        smalls: list[tuple[str, int]] = []
        for label, n in items:
            if n > ROOM_STUDENT_CAP:                         # แตกห้อง (A14)
                k = -(-n // ROOM_STUDENT_CAP)
                base, extra = divmod(n, k)
                for i in range(k):
                    units.append(ExamUnit(f"{label}#{i+1}", base + (1 if i < extra else 0)))
            elif n < SMALL_CLASS:                            # วิชาเล็ก -> พักไว้รวม
                smalls.append((label, n))
            else:
                units.append(ExamUnit(label, n))
        # รวมวิชาเล็กเป็นห้อง ๆ ละ <=65
        smalls.sort(key=lambda e: -e[1])
        bucket_lbl, bucket_n = [], 0
        for label, n in smalls:
            if bucket_n + n > ROOM_STUDENT_CAP and bucket_lbl:
                units.append(ExamUnit("รวมเล็ก: " + ",".join(bucket_lbl), bucket_n))
                bucket_lbl, bucket_n = [], 0
            bucket_lbl.append(label); bucket_n += n
        if bucket_lbl:
            units.append(ExamUnit("รวมเล็ก: " + ",".join(bucket_lbl), bucket_n))
        out[key] = units
    return out
